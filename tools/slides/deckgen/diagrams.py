"""Diagram primitives.

Each function draws into the slide body area and is driven entirely by a spec
dict from the YAML, so translated decks reuse the same geometry.
"""
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import shapes as S
from . import theme as T

# Body area, in inches.
BX, BY, BW, BH = 0.85, 1.55, 11.633, 5.02


def _tone(name):
    return T.TONES.get(name or "neutral", T.TONES["neutral"])


def _I(v):
    return Inches(v)


def _note(slide, text, y=None, color=T.MUTED, size=T.SZ_CAPTION):
    """A caption line under a diagram."""
    y = BY + BH - 0.34 if y is None else y
    return S.textbox(slide, _I(BX), _I(y), _I(BW), _I(0.34), text,
                     size=size, color=color, align=PP_ALIGN.CENTER)


# --------------------------------------------------------------- flow chain
def flow_chain(slide, spec):
    """Horizontal boxes joined by arrows: A → B → C."""
    nodes = spec["nodes"]
    n = len(nodes)
    gap = spec.get("gap", 0.62)
    h = spec.get("height", 1.15)
    y = spec.get("y", BY + (BH - h) / 2 - 0.3)
    w = (BW - gap * (n - 1)) / n

    for i, node in enumerate(nodes):
        x = BX + i * (w + gap)
        fill, line, txt = _tone(node.get("tone"))
        box = S.rect(slide, _I(x), _I(y), _I(w), _I(h), fill=fill, line=line,
                     radius=0.1)
        body = [node["text"]]
        if node.get("note"):
            body.append(node["note"])
        S.label(box, body, size=T.SZ_NODE, color=txt, bold=True)

        if i < n - 1:
            ax = x + w + 0.1
            S.connect(slide, _I(ax), _I(y + h / 2), _I(ax + gap - 0.2),
                      _I(y + h / 2), color=T.DIM, width=1.75)
    if spec.get("caption"):
        _note(slide, spec["caption"], y=y + h + 0.35)


# --------------------------------------------------------------- ladder
def ladder(slide, spec):
    """Ascending rungs, each showing where control sits on a left→right track."""
    rungs = list(reversed(spec["rungs"]))          # highest rung drawn first
    n = len(rungs)
    gap = 0.16
    top = BY + 0.46
    caption_y = BY + BH - 0.32
    room = (caption_y if spec.get("caption") else BY + BH) - top - 0.14
    h = (room - gap * (n - 1)) / n

    axis = spec.get("axis", {})
    S.textbox(slide, _I(BX + 7.35), _I(BY - 0.16), _I(4.05), _I(0.24),
              spec.get("axis_title", "who decides what happens next"),
              size=Pt(9.5), color=T.DIM, align=PP_ALIGN.CENTER, caps=True)
    # Translated labels run longer than English; shrink rather than wrap.
    axis_line = f"{axis.get('left','your code')}   →   {axis.get('right','the model')}"
    axis_size = T.SZ_CAPTION if len(axis_line) <= 32 else Pt(9.0)
    S.textbox(slide, _I(BX + 7.3), _I(BY + 0.08), _I(4.15), _I(0.26), axis_line,
              size=axis_size, color=T.MUTED, align=PP_ALIGN.CENTER, caps=True,
              bold=True)

    for i, rung in enumerate(rungs):
        y = top + i * (h + gap)
        on = rung.get("accent")
        fill, line, txt = _tone("accent" if on else "neutral")

        S.rect(slide, _I(BX), _I(y), _I(BW), _I(h), fill=fill, line=line, radius=0.1)

        chip = S.rect(slide, _I(BX + 0.22), _I(y + (h - 0.5) / 2), _I(0.5), _I(0.5),
                      fill=T.ACCENT if on else T.SURFACE_2, radius=0.12)
        S.label(chip, str(rung["n"]), size=T.SZ_NODE,
                color=T.BG if on else T.MUTED, bold=True)

        S.textbox(slide, _I(BX + 0.92), _I(y + 0.13), _I(6.3), _I(0.3),
                  rung["label"], size=T.SZ_CARD_TITLE, color=T.TEXT, bold=True)
        S.textbox(slide, _I(BX + 0.92), _I(y + 0.44), _I(6.3), _I(h - 0.5),
                  rung.get("note", ""), size=T.SZ_CAPTION, color=T.MUTED,
                  spacing=1.2)

        # control track: dot slides right as autonomy increases
        tx, tw = BX + 7.35, 4.05
        ty = y + h / 2
        S.rect(slide, _I(tx), _I(ty - 0.02), _I(tw), Pt(2.5), fill=T.LINE, radius=0)
        pos = rung.get("control", (rung["n"] - 1) / max(1, len(rungs) - 1))
        dot = S.rect(slide, _I(tx + tw * pos - 0.075), _I(ty - 0.075),
                     _I(0.15), _I(0.15),
                     fill=T.ACCENT if on else T.MUTED, radius=0.075)
        dot.line.fill.background()

    if spec.get("caption"):
        _note(slide, spec["caption"], y=caption_y)


# --------------------------------------------------------------- agent loop
def agent_loop(slide, spec):
    """Linear row with a decision diamond, a looping branch and an exit."""
    y = BY + 0.35
    h = 1.0
    row = spec["nodes"]                      # 2 boxes before the diamond
    dec = spec["decision"]
    branch = spec["branch"]
    ex = spec["exit"]

    w = 2.05
    gap = 0.55
    xs = [BX + i * (w + gap) for i in range(len(row))]
    for x, node in zip(xs, row):
        fill, line, txt = _tone(node.get("tone"))
        box = S.rect(slide, _I(x), _I(y), _I(w), _I(h), fill=fill, line=line, radius=0.1)
        S.label(box, [node["text"]] + ([node["note"]] if node.get("note") else []),
                size=T.SZ_NODE, color=txt, bold=True)
        S.connect(slide, _I(x + w + 0.08), _I(y + h / 2), _I(x + w + gap - 0.08),
                  _I(y + h / 2), color=T.DIM, width=1.75)

    dx = xs[-1] + w + gap
    dw, dh = 2.5, 1.45
    dy = y + h / 2 - dh / 2
    dia = S.diamond(slide, _I(dx), _I(dy), _I(dw), _I(dh))
    S.label(dia, dec["text"], size=T.SZ_NODE_SM, color=T.TEXT, bold=True)

    # exit branch, to the right
    exf, exl, ext = _tone(ex.get("tone", "ok"))
    exx = dx + dw + gap
    exbox = S.rect(slide, _I(exx), _I(y), _I(BX + BW - exx), _I(h),
                   fill=exf, line=exl, radius=0.1)
    S.label(exbox, [ex["text"]] + ([ex["note"]] if ex.get("note") else []),
            size=T.SZ_NODE, color=ext, bold=True)
    S.connect(slide, _I(dx + dw + 0.05), _I(y + h / 2), _I(exx - 0.08),
              _I(y + h / 2), color=T.DIM, width=1.75)
    S.textbox(slide, _I(dx + dw + 0.02), _I(y + h / 2 - 0.34), _I(gap), _I(0.28),
              ex.get("label", "no"), size=T.SZ_CAPTION, color=T.MUTED,
              align=PP_ALIGN.CENTER)

    # looping branch, below
    by = y + h + 1.05
    bnodes = branch["nodes"]
    bw = 2.05
    bxs = [dx + dw / 2 - bw / 2 - i * (bw + gap) for i in range(len(bnodes))]
    S.connect(slide, _I(dx + dw / 2), _I(dy + dh + 0.05), _I(dx + dw / 2),
              _I(by - 0.08), color=T.DIM, width=1.75)
    S.textbox(slide, _I(dx + dw / 2 + 0.1), _I(dy + dh + 0.16), _I(1.4), _I(0.28),
              branch.get("label", "yes"), size=T.SZ_CAPTION, color=T.ACCENT, bold=True)

    for i, (x, node) in enumerate(zip(bxs, bnodes)):
        fill, line, txt = _tone(node.get("tone", "accent"))
        box = S.rect(slide, _I(x), _I(by), _I(bw), _I(h), fill=fill, line=line,
                     radius=0.1)
        S.label(box, [node["text"]] + ([node["note"]] if node.get("note") else []),
                size=T.SZ_NODE, color=txt, bold=True)
        if i < len(bnodes) - 1:
            S.connect(slide, _I(x - 0.08), _I(by + h / 2), _I(x - gap + 0.08),
                      _I(by + h / 2), color=T.DIM, width=1.75, arrow=True)

    # return path: out of the last branch node, left, up, and back into the model
    back_x = bxs[-1]
    llm_x = xs[branch.get("back_to", 1)] + w / 2
    riser = back_x - 0.5
    mid_y = y + h + 0.42
    for x1, y1, x2, y2 in (
        (back_x - 0.06, by + h / 2, riser, by + h / 2),
        (riser, by + h / 2, riser, mid_y),
        (riser, mid_y, llm_x, mid_y),
    ):
        S.connect(slide, _I(x1), _I(y1), _I(x2), _I(y2), color=T.ACCENT,
                  width=1.75, arrow=False)
    S.connect(slide, _I(llm_x), _I(mid_y), _I(llm_x), _I(y + h + 0.05),
              color=T.ACCENT, width=1.75)
    S.textbox(slide, _I(riser + 0.16), _I(mid_y - 0.42), _I(2.4), _I(0.28),
              branch.get("return_label", "loop"), size=T.SZ_CAPTION,
              color=T.ACCENT, bold=True)

    if spec.get("caption"):
        _note(slide, spec["caption"], y=by + h + 0.3)


# --------------------------------------------------------------- stack
def stack(slide, spec):
    """Layered boxes inside an outer container — 'what is on the wire'."""
    items = spec["items"]
    n = len(items)
    outer_w = spec.get("width", 6.4)
    x = BX + spec.get("x_offset", 0.0)
    pad = 0.34
    ih = min(0.82, (BH - 1.2 - pad) / n - 0.13)
    gap = 0.13
    inner_h = n * ih + (n - 1) * gap
    oh = inner_h + pad * 2 + 0.34
    y = BY + max(0.0, (BH - oh) / 2 - 0.25)

    S.rect(slide, _I(x), _I(y), _I(outer_w), _I(oh), fill=T.SURFACE,
           line=T.LINE, radius=0.12)
    S.textbox(slide, _I(x + pad), _I(y + 0.2), _I(outer_w - pad * 2), _I(0.3),
              spec.get("container", "Request"), size=T.SZ_CAPTION,
              color=T.MUTED, caps=True, bold=True)

    for i, item in enumerate(items):
        iy = y + pad + 0.34 + i * (ih + gap)
        fill, line, txt = _tone(item.get("tone"))
        box = S.rect(slide, _I(x + pad), _I(iy), _I(outer_w - pad * 2), _I(ih),
                     fill=fill if item.get("tone") else T.SURFACE_2,
                     line=line if item.get("tone") else None, radius=0.09)
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        S.write(p, item["text"], T.SZ_NODE, txt, bold=True)
        if item.get("note"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            S.write(p2, item["note"], T.SZ_CAPTION, T.MUTED)

    for i, side in enumerate(spec.get("side", [])):
        sy = y + 0.34 + i * 1.15
        S.textbox(slide, _I(x + outer_w + 0.6), _I(sy), _I(BX + BW - x - outer_w - 0.6),
                  _I(0.32), side["text"], size=T.SZ_CARD_TITLE, color=T.TEXT, bold=True)
        S.textbox(slide, _I(x + outer_w + 0.6), _I(sy + 0.36),
                  _I(BX + BW - x - outer_w - 0.6), _I(0.7), side.get("note", ""),
                  size=T.SZ_CAPTION, color=T.MUTED, spacing=1.3)


# --------------------------------------------------------------- barrier
def barrier(slide, spec):
    """Two domains split by a hard line: what the model does vs what your code does."""
    gapw = 0.9
    pw = (BW - gapw) / 2
    y = BY + 0.3
    h = BH - 1.15

    for i, side in enumerate(spec["sides"]):
        x = BX + i * (pw + gapw)
        fill, line, _ = _tone(side.get("tone"))
        S.rect(slide, _I(x), _I(y), _I(pw), _I(h), fill=fill, line=line, radius=0.12)
        S.textbox(slide, _I(x + 0.34), _I(y + 0.3), _I(pw - 0.68), _I(0.34),
                  side["title"], size=T.SZ_LEAD, color=T.TEXT, bold=True)
        S.textbox(slide, _I(x + 0.34), _I(y + 0.7), _I(pw - 0.68), _I(0.3),
                  side.get("subtitle", ""), size=T.SZ_CAPTION, color=T.MUTED, caps=True)
        for j, item in enumerate(side.get("items", [])):
            iy = y + 1.16 + j * 0.62
            box = S.rect(slide, _I(x + 0.34), _I(iy), _I(pw - 0.68), _I(0.5),
                         fill=T.SURFACE_2, radius=0.08)
            p = box.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            S.write(p, item, T.SZ_NODE_SM, T.TEXT)

    mid = BX + pw + gapw / 2
    S.connect(slide, _I(mid), _I(y - 0.05), _I(mid), _I(y + h + 0.05),
              color=T.DANGER, width=1.5, arrow=False, dashed=True)

    if spec.get("caption"):
        _note(slide, spec["caption"], y=y + h + 0.22, color=T.WARN,
              size=T.SZ_CARD_TITLE)


# --------------------------------------------------------------- divergence
def divergence(slide, spec):
    """One input fanning out to several different outputs."""
    src = spec["source"]
    outs = spec["outputs"]
    sw, sh = 2.9, 1.1
    sx, sy = BX + 0.35, BY + (BH - sh) / 2 - 0.35
    box = S.rect(slide, _I(sx), _I(sy), _I(sw), _I(sh), fill=T.SURFACE,
                 line=T.LINE, radius=0.1)
    S.label(box, [src["text"]] + ([src["note"]] if src.get("note") else []),
            size=T.SZ_NODE, color=T.TEXT, bold=True)

    n = len(outs)
    ow = 5.6
    oh = 0.86
    ogap = 0.24
    ox = BX + BW - ow
    total = n * oh + (n - 1) * ogap
    oy0 = BY + (BH - total) / 2 - 0.35

    for i, out in enumerate(outs):
        oy = oy0 + i * (oh + ogap)
        fill, line, txt = _tone(out.get("tone"))
        ob = S.rect(slide, _I(ox), _I(oy), _I(ow), _I(oh), fill=fill, line=line,
                    radius=0.09)
        p = ob.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        S.write(p, out["text"], T.SZ_NODE_SM, txt)
        S.connect(slide, _I(sx + sw + 0.08), _I(sy + sh / 2), _I(ox - 0.08),
                  _I(oy + oh / 2), color=T.DIM, width=1.4)

    if spec.get("caption"):
        _note(slide, spec["caption"], y=oy0 + total + 0.3)


# --------------------------------------------------------------- mapping
def mapping(slide, spec):
    """Left column replaced by right column, pair by pair."""
    pairs = spec["pairs"]
    n = len(pairs)
    h = min(0.72, (BH - 1.15) / n - 0.14)
    gap = 0.14
    cw = 4.75
    lx = BX
    rx = BX + BW - cw
    y0 = BY + 0.62

    S.textbox(slide, _I(lx), _I(BY + 0.12), _I(cw), _I(0.3),
              spec.get("left_title", "What breaks"), size=T.SZ_CAPTION,
              color=T.DANGER, bold=True, caps=True)
    S.textbox(slide, _I(rx), _I(BY + 0.12), _I(cw), _I(0.3),
              spec.get("right_title", "What replaces it"), size=T.SZ_CAPTION,
              color=T.OK, bold=True, caps=True)

    for i, (left, right) in enumerate((p["from"], p["to"]) for p in pairs):
        y = y0 + i * (h + gap)
        lb = S.rect(slide, _I(lx), _I(y), _I(cw), _I(h), fill=T.DANGER_SOFT,
                    line=T.DANGER, radius=0.09)
        p = lb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        S.write(p, left, T.SZ_NODE_SM, T.TEXT)

        rb = S.rect(slide, _I(rx), _I(y), _I(cw), _I(h), fill=T.OK_SOFT,
                    line=T.OK, radius=0.09)
        p = rb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        S.write(p, right, T.SZ_NODE_SM, T.TEXT)

        S.connect(slide, _I(lx + cw + 0.12), _I(y + h / 2), _I(rx - 0.12),
                  _I(y + h / 2), color=T.DIM, width=1.5)


# --------------------------------------------------------------- flowchart
def flowchart(slide, spec):
    """A vertical question chain, each question branching out to an outcome."""
    steps = spec["steps"]
    n = len(steps)
    final_h = 0.66 if spec.get("final") else 0.0
    qw = 4.3
    # Spread the chain over the whole body rather than bunching it at the top.
    qh = min(0.95, (BH - 0.3 - final_h) / (n + 0.55 * bool(final_h)) * 0.74)
    gap = ((BH - 0.3 - final_h) - n * qh) / max(1, n - 0.45)
    x = BX + 0.15
    y0 = BY + 0.1
    ow = 5.3
    ox = BX + BW - ow

    for i, step in enumerate(steps):
        y = y0 + i * (qh + gap)
        d = S.rect(slide, _I(x), _I(y), _I(qw), _I(qh), fill=T.SURFACE,
                   line=T.ACCENT, radius=0.4)
        S.label(d, step["q"], size=T.SZ_NODE_SM, color=T.TEXT, bold=True)

        br = step["branch"]
        fill, line, txt = _tone(br.get("tone", "ok"))
        ob = S.rect(slide, _I(ox), _I(y), _I(ow), _I(qh), fill=fill, line=line,
                    radius=0.09)
        p = ob.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        S.write(p, br["text"], T.SZ_NODE_SM, txt)

        S.connect(slide, _I(x + qw + 0.1), _I(y + qh / 2), _I(ox - 0.1),
                  _I(y + qh / 2), color=T.DIM, width=1.4)
        S.textbox(slide, _I(x + qw + 0.2), _I(y + qh / 2 - 0.3), _I(0.7), _I(0.26),
                  br.get("label", "yes"), size=T.SZ_CAPTION, color=T.MUTED)

        if i < n - 1:
            S.connect(slide, _I(x + qw / 2), _I(y + qh + 0.05), _I(x + qw / 2),
                      _I(y + qh + gap - 0.05), color=T.DIM, width=1.5)
            S.textbox(slide, _I(x + qw / 2 + 0.14), _I(y + qh + 0.08), _I(0.8),
                      _I(0.26), step.get("next_label", "no"), size=T.SZ_CAPTION,
                      color=T.MUTED)

    if spec.get("final"):
        f = spec["final"]
        fy = y0 + (n - 1) * (qh + gap) + qh + gap * 0.45
        fill, line, txt = _tone(f.get("tone", "accent"))
        fb = S.rect(slide, _I(x), _I(fy), _I(qw), _I(final_h - 0.06), fill=fill,
                    line=line, radius=0.09)
        S.label(fb, f["text"], size=T.SZ_NODE_SM, color=txt, bold=True)
        S.connect(slide, _I(x + qw / 2), _I(fy - gap * 0.5), _I(x + qw / 2),
                  _I(fy - 0.05), color=T.DIM, width=1.5)


# --------------------------------------------------------------- matrix
def matrix(slide, spec):
    """Grid of short verdicts — cases down the side, questions across the top."""
    cols = spec["columns"]
    rows = spec["rows"]
    label_w = spec.get("label_width", 3.5)
    nc = len(cols)
    cw = (BW - label_w) / nc
    rh = min(0.78, (BH - 1.1) / (len(rows) + 1))
    y0 = BY + 0.52          # headers may wrap to two lines

    for j, col in enumerate(cols):
        S.textbox(slide, _I(BX + label_w + j * cw), _I(y0 - 0.5), _I(cw - 0.1),
                  _I(0.46), col, size=T.SZ_CAPTION, color=T.MUTED, bold=True,
                  align=PP_ALIGN.CENTER, caps=True, anchor=MSO_ANCHOR.BOTTOM)

    for i, row in enumerate(rows):
        y = y0 + i * (rh + 0.11)
        lb = S.rect(slide, _I(BX), _I(y), _I(label_w - 0.14), _I(rh),
                    fill=T.SURFACE, line=T.LINE, radius=0.09)
        S.label(lb, [row["label"]] + ([row["note"]] if row.get("note") else []),
                size=T.SZ_NODE_SM, color=T.TEXT, bold=True, align=PP_ALIGN.LEFT)

        for j, cell in enumerate(row["cells"]):
            if isinstance(cell, str):
                cell = {"text": cell}
            fill, line, txt = _tone(cell.get("tone"))
            last = j == len(row["cells"]) - 1
            cb = S.rect(slide, _I(BX + label_w + j * cw), _I(y), _I(cw - 0.14),
                        _I(rh), fill=fill, line=line, radius=0.09)
            S.label(cb, cell["text"], size=T.SZ_NODE_SM, color=txt, bold=last)


# --------------------------------------------------------------- quadrant
def quadrant(slide, spec):
    """Four cards in a 2×2 — used where a concept genuinely has four parts."""
    cells = spec["cells"]
    gap = 0.28
    cw = (BW - gap) / 2
    ch = (BH - 0.85 - gap) / 2
    y0 = BY + 0.18

    for i, cell in enumerate(cells[:4]):
        x = BX + (i % 2) * (cw + gap)
        y = y0 + (i // 2) * (ch + gap)
        fill, line, txt = _tone(cell.get("tone"))
        box = S.rect(slide, _I(x), _I(y), _I(cw), _I(ch), fill=fill, line=line,
                     radius=0.12)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Emu(160000)
        tf.margin_top = Emu(130000)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        S.write(p, cell["title"], T.SZ_LEAD, txt, bold=True)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        p2.line_spacing = 1.3
        p2.space_before = Pt(8)
        S.write(p2, cell["text"], T.SZ_CARD_BODY, T.MUTED)


# --------------------------------------------------------------- nested compare
def nested_compare(slide, spec):
    """Two architectures side by side, each as nested containers.

    Group heights are solved against the space available so a panel with more
    boxes shrinks to fit rather than running over the verdict line.
    """
    panels = spec["panels"]
    gapw = 0.85
    pw = (BW - gapw) / 2
    top = BY + 0.04
    title_h = 0.70          # panel titles may run to two lines
    verdict_h = 0.66 if any(p.get("verdict") for p in panels) else 0.0
    avail = BH - 0.08 - title_h - verdict_h

    head, step, pad, div_h, ggap = 0.34, 0.42, 0.16, 0.56, 0.18

    def needed(panel):
        total = 0.0
        for group in panel["groups"]:
            total += (div_h if group.get("divider")
                      else head + len(group.get("items", [])) * step + pad)
            total += ggap
        return max(0.0, total - ggap)

    worst = max(needed(p) for p in panels)
    if worst > avail:
        k = avail / worst
        head, step, pad, div_h, ggap = (v * k for v in (head, step, pad, div_h, ggap))

    for i, panel in enumerate(panels):
        x = BX + i * (pw + gapw)
        good = panel.get("verdict_tone") == "ok"
        S.textbox(slide, _I(x), _I(top), _I(pw), _I(title_h), panel["title"],
                  size=T.SZ_CARD_TITLE, color=T.OK if good else T.WARN, bold=True,
                  spacing=1.2)

        gy = top + title_h
        for group in panel["groups"]:
            if group.get("divider"):
                S.connect(slide, _I(x + 0.1), _I(gy + div_h * 0.34),
                          _I(x + pw - 0.1), _I(gy + div_h * 0.34), color=T.DANGER,
                          width=1.4, arrow=False, dashed=True)
                S.textbox(slide, _I(x), _I(gy + div_h * 0.44), _I(pw), _I(0.28),
                          group["divider"], size=T.SZ_CAPTION, color=T.DANGER,
                          align=PP_ALIGN.CENTER, caps=True, bold=True)
                gy += div_h + ggap
                continue

            items = group.get("items", [])
            gh = head + len(items) * step + pad
            fill, line, _ = _tone(group.get("tone"))
            S.rect(slide, _I(x), _I(gy), _I(pw), _I(gh), fill=fill, line=line,
                   radius=0.11)
            S.textbox(slide, _I(x + 0.24), _I(gy + 0.11), _I(pw - 0.48), _I(0.26),
                      group["title"], size=T.SZ_CAPTION, color=T.MUTED,
                      caps=True, bold=True)
            for j, item in enumerate(items):
                if isinstance(item, str):
                    item = {"text": item}
                ifill, iline, itxt = _tone(item.get("tone"))
                ib = S.rect(slide, _I(x + 0.24), _I(gy + head + j * step),
                            _I(pw - 0.48), _I(step - 0.06),
                            fill=ifill if item.get("tone") else T.SURFACE_2,
                            line=iline if item.get("tone") else None, radius=0.07)
                p = ib.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                S.write(p, item["text"], T.SZ_NODE_SM, itxt,
                        bold=bool(item.get("tone")))
            gy += gh + ggap

        if panel.get("verdict"):
            S.textbox(slide, _I(x), _I(BY + BH - verdict_h + 0.06), _I(pw),
                      _I(verdict_h - 0.1), panel["verdict"], size=T.SZ_CAPTION,
                      color=T.OK if good else T.WARN, bold=True, spacing=1.25)


# --------------------------------------------------------------- ceiling
def ceiling(slide, spec):
    """A horizontal budget bar running into a hard limit."""
    segs = spec["segments"]
    y = BY + 0.75
    h = 1.35
    total = sum(s["pct"] for s in segs)
    x = BX
    bar_w = BW - 1.4          # room for the limit rule and its label

    for seg in segs:
        w = bar_w * seg["pct"] / total
        fill, line, txt = _tone(seg.get("tone"))
        box = S.rect(slide, _I(x), _I(y), _I(w - 0.04), _I(h),
                     fill=fill, line=line, radius=0.08)
        S.label(box, [seg["text"], seg.get("note", "")] if seg.get("note")
                else [seg["text"]], size=T.SZ_NODE_SM, color=txt, bold=True)
        x += w

    limit_x = BX + bar_w + 0.08
    S.connect(slide, _I(limit_x), _I(y - 0.3), _I(limit_x), _I(y + h + 0.3),
              color=T.DANGER, width=2.5, arrow=False)
    S.textbox(slide, _I(limit_x + 0.14), _I(y + h / 2 - 0.34), _I(1.15), _I(0.8),
              spec.get("limit", "Hard\nlimit"), size=T.SZ_CAPTION, color=T.DANGER,
              bold=True, spacing=1.25)

    if spec.get("caption"):
        _note(slide, spec["caption"], y=y + h + 0.5, color=T.WARN,
              size=T.SZ_CARD_TITLE)


def layered(slide, spec):
    """Horizontal bands, each optionally holding boxes — architecture pictures.

    Band heights come from relative weights, so a container band can dominate
    while a cross-cutting bar stays a strip.
    """
    bands = spec["bands"]
    gap = spec.get("gap", 0.22)
    caption = spec.get("caption")
    bottom = (BY + BH - 0.34) if caption else (BY + BH)
    top = BY + 0.1
    total_w = sum(b.get("weight", 1) for b in bands)
    room = bottom - top - gap * (len(bands) - 1)

    y = top
    for band in bands:
        h = room * band.get("weight", 1) / total_w
        fill, line, txt = _tone(band.get("tone"))
        S.rect(slide, _I(BX), _I(y), _I(BW), _I(h), fill=fill, line=line,
               radius=0.12)
        boxes = band.get("boxes", [])
        if boxes:
            S.textbox(slide, _I(BX + 0.3), _I(y + 0.13), _I(BW - 0.6), _I(0.26),
                      band["title"], size=T.SZ_CAPTION, color=T.MUTED, caps=True,
                      bold=True)
            if band.get("note"):
                S.textbox(slide, _I(BX + 0.3), _I(y + 0.4), _I(BW - 0.6), _I(0.26),
                          band["note"], size=T.SZ_CAPTION, color=T.DIM)
        else:
            # A strip with no boxes: title and note share one centred line.
            ty = y + (h - 0.26) / 2
            S.textbox(slide, _I(BX + 0.3), _I(ty), _I(2.6), _I(0.26),
                      band["title"], size=T.SZ_CAPTION, color=T.MUTED, caps=True,
                      bold=True)
            if band.get("note"):
                S.textbox(slide, _I(BX + 3.0), _I(ty), _I(BW - 3.3), _I(0.26),
                          band["note"], size=T.SZ_CAPTION, color=T.DIM)

        if boxes:
            n = len(boxes)
            pad = 0.3
            bgap = 0.2
            head = 0.72 if band.get("note") else 0.48
            bw = (BW - pad * 2 - bgap * (n - 1)) / n
            bh = h - head - pad
            for j, box in enumerate(boxes):
                if isinstance(box, str):
                    box = {"text": box}
                bfill, bline, btxt = _tone(box.get("tone"))
                shp = S.rect(slide, _I(BX + pad + j * (bw + bgap)), _I(y + head),
                             _I(bw), _I(bh),
                             fill=bfill if box.get("tone") else T.SURFACE_2,
                             line=bline if box.get("tone") else None, radius=0.1)
                S.label(shp, [box["text"]] + ([box["note"]] if box.get("note") else []),
                        size=T.SZ_NODE, color=btxt, bold=True)
        y += h + gap

    if caption:
        _note(slide, caption, y=BY + BH - 0.3)


RENDERERS = {
    "flow_chain": flow_chain,
    "layered": layered,
    "ladder": ladder,
    "agent_loop": agent_loop,
    "stack": stack,
    "barrier": barrier,
    "divergence": divergence,
    "mapping": mapping,
    "flowchart": flowchart,
    "matrix": matrix,
    "quadrant": quadrant,
    "nested_compare": nested_compare,
    "ceiling": ceiling,
}
