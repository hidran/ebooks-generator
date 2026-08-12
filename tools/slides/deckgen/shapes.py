"""Low-level shape and text primitives.

Everything drawn by the engine goes through here, so the visual rules (inline
`code` styling, arrowheads, corner radii) are enforced in exactly one place.
"""
import re

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import theme as T

# `code` → mono, **bold** → bold. The lesson scripts lean on both heavily.
_TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def rich_runs(text):
    """Split a string into (text, kind) pairs where kind is plain|mono|bold."""
    out = []
    for part in _TOKEN.split(str(text)):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) > 1:
            out.append((part[1:-1], "mono"))
        elif part.startswith("**") and part.endswith("**") and len(part) > 3:
            out.append((part[2:-2], "bold"))
        else:
            out.append((part, "plain"))
    return out or [("", "plain")]


def write(para, text, size, color, bold=False, font=T.FONT, italic=False):
    """Fill a paragraph with rich runs, honouring inline markup."""
    for chunk, kind in rich_runs(text):
        run = para.add_run()
        run.text = chunk
        f = run.font
        f.size = Pt(size.pt * 0.94) if kind == "mono" else size
        f.name = T.MONO if kind == "mono" else font
        f.bold = bold or kind == "bold"
        f.italic = italic
        f.color.rgb = color
    return para


def textbox(slide, x, y, w, h, text="", size=T.SZ_BODY, color=T.TEXT, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=T.FONT,
            spacing=1.15, italic=False, caps=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        if caps:
            line = line.upper()
        write(para, line, size, color, bold=bold, font=font, italic=italic)
        if caps:
            for run in para.runs:
                run.font._rPr.set("spc", "160")
    return box


def _no_line(shape):
    shape.line.fill.background()


def rect(slide, x, y, w, h, fill=T.SURFACE, line=None, radius=None, width=1.0,
         shape_type=None):
    """A rounded rectangle by default; radius is in inches."""
    st = shape_type or (MSO_SHAPE.ROUNDED_RECTANGLE if radius is None or radius > 0
                        else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(st, x, y, w, h)
    shp.shadow.inherit = False

    if st == MSO_SHAPE.ROUNDED_RECTANGLE:
        r = 0.09 if radius is None else radius
        smallest = min(Emu(w).inches, Emu(h).inches)
        try:
            shp.adjustments[0] = max(0.0, min(0.5, r / smallest if smallest else 0))
        except IndexError:
            pass

    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill

    if line is None:
        _no_line(shp)
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(width)

    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(60000)
    tf.margin_top = tf.margin_bottom = Emu(40000)
    return shp


def label(shp, lines, size=T.SZ_NODE, color=T.TEXT, bold=False, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE, spacing=1.1):
    """Write centred text into an existing autoshape."""
    tf = shp.text_frame
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        strong = bold if i == 0 else False
        col = color if i == 0 else T.MUTED
        sz = size if i == 0 else Pt(size.pt * 0.82)
        write(para, line, sz, col, bold=strong)
    return shp


def _arrow(shp, head=False, tail=True):
    ln = shp.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if not on:
            continue
        el = ln.makeelement(qn(tag), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(el)


def connect(slide, x1, y1, x2, y2, color=T.LINE, width=1.5, arrow=True,
            head=False, kind=MSO_CONNECTOR.STRAIGHT, dashed=False):
    c = slide.shapes.add_connector(kind, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = 4  # msoLineDash
    _arrow(c, head=head, tail=arrow)
    return c


def elbow(slide, x1, y1, x2, y2, **kw):
    kw.setdefault("kind", MSO_CONNECTOR.ELBOW)
    return connect(slide, x1, y1, x2, y2, **kw)


def chevron(slide, x, y, w, h, fill=T.SURFACE, line=T.LINE):
    shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, y, w, h)
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        _no_line(shp)
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.0)
    shp.text_frame.word_wrap = True
    return shp


def diamond(slide, x, y, w, h, fill=T.ACCENT_SOFT, line=T.ACCENT):
    shp = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, x, y, w, h)
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.25)
    shp.text_frame.word_wrap = True
    return shp


def pill(slide, x, y, w, h, text, fill=T.SURFACE_2, line=None, color=T.TEXT,
         size=T.SZ_NODE_SM, bold=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.5
    except IndexError:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        _no_line(shp)
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.0)
    shp.text_frame.margin_left = shp.text_frame.margin_right = Emu(50000)
    shp.text_frame.margin_top = shp.text_frame.margin_bottom = 0
    return label(shp, text, size=size, color=color, bold=bold)


def rule(slide, x, y, w, color=T.ACCENT, weight=3.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(weight))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    _no_line(shp)
    return shp
