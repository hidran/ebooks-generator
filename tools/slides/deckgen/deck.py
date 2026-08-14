"""The Deck wrapper: slide chrome shared by every layout."""
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import shapes as S
from . import theme as T


class Deck:
    def __init__(self, title, subtitle="", module="", lessons=None):
        self.prs = Presentation()
        self.prs.slide_width = T.SLIDE_W
        self.prs.slide_height = T.SLIDE_H
        self.title = title
        self.subtitle = subtitle
        self.module = module
        self.lessons = lessons or []      # ordered lesson ids, e.g. ["1.1", "1.2"]
        self.current = None               # footer context: (id, title)

    # ------------------------------------------------------------ chrome
    def blank(self, footer=True):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = T.BG
        if footer:
            self._footer(slide)
        return slide

    def _footer(self, slide):
        if not self.current:
            return
        lid, ltitle = self.current
        S.textbox(slide, T.MARGIN_X, T.FOOTER_Y, Inches(8.5), Inches(0.3),
                  f"{lid} · {ltitle}", size=T.SZ_FOOTER, color=T.DIM)
        S.textbox(slide, Inches(9.6), T.FOOTER_Y, Inches(2.9), Inches(0.3),
                  self.title, size=T.SZ_FOOTER, color=T.DIM, align=PP_ALIGN.RIGHT)
        self._progress(slide, lid)

    def _progress(self, slide, lid):
        """A segmented strip: which lesson of the module is on screen."""
        if len(self.lessons) < 2:
            return
        total = len(self.lessons)
        span = Emu(T.CONTENT_W)
        gap = Emu(Inches(0.06))
        seg = int((span - gap * (total - 1)) / total)
        for i, lesson in enumerate(self.lessons):
            x = Emu(T.MARGIN_X) + i * (seg + gap)
            on = lesson == lid
            S.rect(slide, Emu(x), T.PROGRESS_Y, Emu(seg), Pt(2.5),
                   fill=T.ACCENT if on else T.LINE, radius=0)

    def heading(self, slide, title, kicker=None):
        """Kicker in accent, title in white. No underline rule — at this size it
        reads as a typo rather than as structure."""
        y = T.TITLE_Y
        if kicker:
            S.textbox(slide, T.MARGIN_X, Inches(0.44), T.CONTENT_W, Inches(0.26),
                      kicker, size=T.SZ_KICKER, color=T.ACCENT, bold=True, caps=True)
            y = Inches(0.76)
        # ~52 characters fit one line at the full size. Step down rather than
        # letting a wrapped title crowd the body area below it.
        n = len(str(title))
        size = T.SZ_TITLE if n <= 50 else (Pt(24) if n <= 105 else Pt(20))
        S.textbox(slide, T.MARGIN_X, y, T.CONTENT_W, Inches(0.72), title,
                  size=size, color=T.TEXT, bold=True, spacing=1.04)

    # ------------------------------------------------------------ notes
    def speaker_notes(self, slide, text):
        """Delivery cues for the presenter view — what to say on this slide."""
        if not text:
            return None
        frame = slide.notes_slide.notes_text_frame
        lines = str(text).rstrip("\n").split("\n")
        frame.text = lines[0]
        for line in lines[1:]:
            frame.add_paragraph().text = line
        for para in frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(13)
                run.font.name = T.FONT
        return frame

    # ------------------------------------------------------------ output
    def _link_notes_master(self):
        """Declare the notes master in presentation.xml.

        Adding a notes slide makes python-pptx create the notesMaster part and
        relate presentation.xml to it, but it never writes the matching
        <p:notesMasterIdLst>. That leaves the relationship dangling: PowerPoint
        repairs it silently, Keynote refuses to open the file at all. Write the
        element ourselves, in schema order (right after sldMasterIdLst).
        """
        prs_elm = self.prs._element
        if prs_elm.find(qn("p:notesMasterIdLst")) is not None:
            return
        rel = next((r for r in self.prs.part.rels.values()
                    if r.reltype == RT.NOTES_MASTER), None)
        if rel is None:                       # a deck with no speaker notes
            return
        lst = prs_elm.makeelement(qn("p:notesMasterIdLst"), {})
        lst.append(lst.makeelement(qn("p:notesMasterId"), {qn("r:id"): rel.rId}))
        masters = prs_elm.find(qn("p:sldMasterIdLst"))
        if masters is None:
            prs_elm.insert(0, lst)
        else:
            masters.addnext(lst)

    def _declare_aspect(self):
        """The stock template says 4:3; these decks are 16:9. PowerPoint ignores
        the mismatch, but leaving a false claim in the file invites trouble from
        stricter readers."""
        sld_sz = self.prs._element.find(qn("p:sldSz"))
        if sld_sz is not None:
            sld_sz.set("type", "screen16x9")

    def save(self, path):
        self._link_notes_master()
        self._declare_aspect()
        self.prs.save(str(path))
        return path

    @property
    def slide_count(self):
        return len(self.prs.slides._sldIdLst)
