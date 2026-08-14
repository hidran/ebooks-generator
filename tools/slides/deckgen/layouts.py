"""Text-driven slide layouts."""
import math

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from . import shapes as S
from . import theme as T
from .diagrams import BH, BW, BX, BY, _I


# Measured against rendered Inter Bold: a character is ~0.0075in per point of
# font size, and a line at spacing 1.2 occupies ~0.020in per point.
CHAR_IN_PER_PT = 0.0075
LINE_IN_PER_PT = 0.0200


def _fit_size(text, width_in, avail_in, max_pt=34, min_pt=16):
    """Largest size at which `text` still fits `avail_in`, honouring newlines."""
    for pt in range(int(max_pt), int(min_pt) - 1, -1):
        chars_per_line = max(8, int(width_in / (CHAR_IN_PER_PT * pt)))
        lines = sum(max(1, math.ceil(len(para) / chars_per_line))
                    for para in str(text).split("\n"))
        if lines * LINE_IN_PER_PT * pt <= avail_in:
            return pt
    return min_pt


# --------------------------------------------------------------- covers
def cover(deck, slide, spec):
    S.rect(slide, _I(0), _I(0), _I(0.14), T.SLIDE_H, fill=T.ACCENT, radius=0)
    S.textbox(slide, _I(1.35), _I(2.05), _I(10.2), _I(0.34),
              spec.get("kicker", deck.module), size=T.SZ_KICKER, color=T.ACCENT,
              bold=True, caps=True)
    S.textbox(slide, _I(1.35), _I(2.55), _I(10.2), _I(1.5), spec["title"],
              size=T.SZ_DECK_TITLE, color=T.TEXT, bold=True, spacing=1.05)
    S.rule(slide, _I(1.35), _I(4.22), _I(1.6))
    S.textbox(slide, _I(1.35), _I(4.55), _I(9.4), _I(0.8), spec.get("subtitle", ""),
              size=T.SZ_LEAD, color=T.MUTED, spacing=1.3)
    S.textbox(slide, _I(1.35), _I(5.85), _I(9.4), _I(0.4), spec.get("meta", ""),
              size=T.SZ_CAPTION, color=T.DIM, caps=True)


def divider(deck, slide, spec):
    """Lesson opener: number and title.

    Run time is deliberately absent. It is a cue for whoever is recording, not
    information the viewer needs, and on screen it only invites clock-watching.
    build() moves it into the speaker notes instead.
    """
    S.rect(slide, _I(0), _I(0), _I(0.14), T.SLIDE_H, fill=T.ACCENT, radius=0)
    S.textbox(slide, _I(1.35), _I(1.95), _I(3.0), _I(1.4), spec["number"],
              size=T.SZ_DIVIDER_NUM, color=T.ACCENT, bold=True, spacing=0.9)
    S.textbox(slide, _I(1.35), _I(3.42), _I(10.0), _I(1.3), spec["title"],
              size=T.SZ_DIVIDER_TITLE, color=T.TEXT, bold=True, spacing=1.05)
    S.textbox(slide, _I(1.35), _I(4.92), _I(9.0), _I(0.34), spec.get("type") or "",
              size=T.SZ_CAPTION, color=T.DIM, caps=True)


def agenda(deck, slide, spec):
    deck.heading(slide, spec.get("title", "In this module"), spec.get("kicker"))
    items = spec["items"]
    half = (len(items) + 1) // 2
    cols = [items[:half], items[half:]]
    cw = (BW - 0.5) / 2
    for c, column in enumerate(cols):
        for i, item in enumerate(column):
            y = BY + 0.1 + i * 0.78
            x = BX + c * (cw + 0.5)
            chip = S.rect(slide, _I(x), _I(y), _I(0.72), _I(0.5),
                          fill=T.SURFACE_2, radius=0.1)
            S.label(chip, item["id"], size=T.SZ_NODE_SM, color=T.ACCENT, bold=True)
            S.textbox(slide, _I(x + 0.9), _I(y + 0.03), _I(cw - 1.0), _I(0.3),
                      item["title"], size=T.SZ_CARD_TITLE, color=T.TEXT, bold=True)
            S.textbox(slide, _I(x + 0.9), _I(y + 0.3), _I(cw - 1.0), _I(0.26),
                      item.get("note", ""), size=T.SZ_CAPTION, color=T.MUTED)


# --------------------------------------------------------------- text
def objectives(deck, slide, spec):
    deck.heading(slide, spec.get("title", "What you will be able to do"),
                 spec.get("kicker", "Learning objectives"))
    points = spec.get("points", [])
    # With no sub-points the slide is mostly empty; sit the box on the optical
    # centre instead of leaving it stranded at the top.
    y0 = BY + 0.15 if points else BY + 1.1
    S.rect(slide, _I(BX), _I(y0), _I(BW), _I(1.9), fill=T.ACCENT_SOFT,
           line=T.ACCENT, radius=0.14)
    S.textbox(slide, _I(BX + 0.55), _I(y0 + 0.38), _I(BW - 1.1), _I(1.3),
              spec["text"], size=T.SZ_LEAD, color=T.TEXT, spacing=1.35)
    for i, item in enumerate(points):
        y = y0 + 2.28 + i * 0.62
        S.rect(slide, _I(BX + 0.06), _I(y + 0.16), _I(0.16), _I(0.16),
               fill=T.ACCENT, radius=0.08)
        S.textbox(slide, _I(BX + 0.46), _I(y), _I(BW - 0.46), _I(0.55), item,
                  size=T.SZ_BODY, color=T.MUTED, spacing=1.3)


def bullets(deck, slide, spec):
    deck.heading(slide, spec["title"], spec.get("kicker"))
    items = spec["items"]
    y = BY + 0.1
    if spec.get("lead"):
        S.textbox(slide, _I(BX), _I(y), _I(BW), _I(0.75), spec["lead"],
                  size=T.SZ_LEAD, color=T.TEXT, spacing=1.3)
        y += 0.95
    step = min(0.92, (BY + BH - 0.25 - y) / max(1, len(items)))
    for i, item in enumerate(items):
        if isinstance(item, str):
            item = {"text": item}
        iy = y + i * step
        marker = item.get("marker")
        if marker:
            chip = S.rect(slide, _I(BX), _I(iy + 0.02), _I(0.46), _I(0.4),
                          fill=T.SURFACE_2, radius=0.09)
            S.label(chip, marker, size=T.SZ_NODE_SM, color=T.ACCENT, bold=True)
        else:
            S.rect(slide, _I(BX + 0.06), _I(iy + 0.16), _I(0.15), _I(0.15),
                   fill=T.ACCENT, radius=0.075)
        tx = BX + (0.66 if marker else 0.46)
        S.textbox(slide, _I(tx), _I(iy), _I(BW - (tx - BX)), _I(step),
                  item["text"], size=T.SZ_BODY, color=T.TEXT, spacing=1.25)
        if item.get("note"):
            S.textbox(slide, _I(tx), _I(iy + 0.34), _I(BW - (tx - BX)), _I(step - 0.3),
                      item["note"], size=T.SZ_CARD_BODY, color=T.MUTED, spacing=1.25)


def lead(deck, slide, spec):
    """One sentence that deserves the whole slide.

    The statement is set large, so its size is solved against the space the note
    leaves rather than fixed — a long lead shrinks instead of running off.
    """
    if spec.get("title"):
        deck.heading(slide, spec["title"], spec.get("kicker"))
        y = BY + 0.30
    else:
        y = BY + 0.85
    bottom = BY + BH

    note = spec.get("note")
    note_h = 0.0
    note_size = T.SZ_BODY
    if note:
        if len(note) > 210:
            note_size = T.SZ_CARD_BODY
        chars = 128 if note_size == T.SZ_CARD_BODY else 96
        lines = max(1, math.ceil(len(note) / chars))
        note_h = (note_size.pt / 72 * 1.35) * lines + 0.22

    S.rule(slide, _I(BX), _I(y), _I(1.35))
    ty = y + 0.40
    avail = bottom - note_h - ty - 0.18

    text = str(spec["text"])
    size = spec.get("size") or _fit_size(text, BW - 1.0, max(0.6, avail))
    S.textbox(slide, _I(BX), _I(ty), _I(BW - 1.0), _I(max(0.6, avail)), text,
              size=Pt(size), color=T.TEXT, bold=True, spacing=1.2)

    if note:
        S.textbox(slide, _I(BX), _I(bottom - note_h), _I(BW - 1.0), _I(note_h),
                  note, size=note_size, color=T.MUTED, spacing=1.35)


def cards(deck, slide, spec):
    deck.heading(slide, spec["title"], spec.get("kicker"))
    items = spec["items"]
    n = len(items)
    per_row = spec.get("columns", min(n, 4 if n % 4 == 0 or n > 3 else n))
    rows = (n + per_row - 1) // per_row
    gap = 0.28
    cw = (BW - gap * (per_row - 1)) / per_row
    avail = BH - 0.55
    ch = min(spec.get("height", 2.7), (avail - gap * (rows - 1)) / rows)
    # A single row stranded at the top reads as an unfinished slide; centre it.
    y0 = BY + ((avail - ch) / 2 if rows == 1 else 0.12)

    # Card text shrinks to fit the box rather than running out of the bottom of
    # it. Longer languages wrap where English did not, and a fixed size clips.
    text_w = cw - 0.36

    def _needed(item, title_pt, body_pt):
        def wrapped(text, pt):
            per_line = max(6, int(text_w / (CHAR_IN_PER_PT * pt)))
            return sum(max(1, math.ceil(len(seg) / per_line))
                       for seg in str(text).split("\n"))
        h = 0.25  # top margin plus bottom breathing room
        if item.get("marker"):
            h += wrapped(item["marker"], T.SZ_CAPTION.pt) * T.SZ_CAPTION.pt * 0.020
            h += 0.055
        h += wrapped(item["title"], title_pt) * title_pt * 0.020
        if item.get("text"):
            h += 0.097
            h += wrapped(item["text"], body_pt) * body_pt * 0.0226
        return h

    body_pt = T.SZ_CARD_BODY.pt
    title_pt = T.SZ_CARD_TITLE.pt
    while body_pt > 8.0:
        if max(_needed(it, title_pt, body_pt) for it in items) <= ch:
            break
        body_pt -= 0.5
        title_pt = max(11.0, title_pt - 0.5)
    body_size, title_size = Pt(body_pt), Pt(title_pt)

    for i, item in enumerate(items):
        x = BX + (i % per_row) * (cw + gap)
        y = y0 + (i // per_row) * (ch + gap)
        fill, line, txt = T.TONES.get(item.get("tone", "neutral"), T.TONES["neutral"])
        box = S.rect(slide, _I(x), _I(y), _I(cw), _I(ch), fill=fill, line=line,
                     radius=0.12)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Emu(150000)
        tf.margin_top = Emu(120000)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        if item.get("marker"):
            S.write(p, item["marker"], T.SZ_CAPTION, T.ACCENT, bold=True)
            p = tf.add_paragraph()
            p.space_before = Pt(4)
            p.alignment = PP_ALIGN.LEFT
        S.write(p, item["title"], title_size, txt, bold=True)
        if item.get("text"):
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            p2.line_spacing = 1.3
            p2.space_before = Pt(7)
            S.write(p2, item["text"], body_size, T.MUTED)


def table(deck, slide, spec):
    deck.heading(slide, spec["title"], spec.get("kicker"))
    cols = spec["columns"]
    rows = spec["rows"]
    widths = spec.get("widths") or [1] * len(cols)
    scale = BW / sum(widths)
    y0 = BY + 0.12

    # Rows grow to fit a cell that wraps. A fixed row height silently clips the
    # second line, and translations wrap where English did not.
    size = T.SZ_TABLE
    line_h = size.pt * LINE_IN_PER_PT

    def _lines(cells):
        n = 1
        for j, cell in enumerate(cells):
            avail_w = widths[j] * scale - 0.28
            per_line = max(6, int(avail_w / (CHAR_IN_PER_PT * size.pt)))
            n = max(n, math.ceil(len(str(cell)) / per_line))
        return n

    head_h = 0.46
    body = [row["cells"] if isinstance(row, dict) else row for row in rows]
    heights = [max(0.46, _lines(cells) * line_h + 0.24) for cells in body]

    avail = BH - 0.12 - (0.5 if spec.get("caption") else 0.1)
    if head_h + sum(heights) > avail:
        k = (avail - head_h) / sum(heights)
        heights = [h * k for h in heights]

    x = BX
    for j, col in enumerate(cols):
        w = widths[j] * scale
        S.textbox(slide, _I(x + 0.14), _I(y0 + 0.08), _I(w - 0.2), _I(0.3), col,
                  size=T.SZ_CAPTION, color=T.MUTED, bold=True, caps=True)
        x += w
    S.rect(slide, _I(BX), _I(y0 + head_h - 0.03), _I(BW), Pt(1.2), fill=T.LINE,
           radius=0)

    y = y0 + head_h
    for i, (row, cells, rh) in enumerate(zip(rows, body, heights)):
        emphasis = isinstance(row, dict) and row.get("tone")
        if emphasis:
            fill, line, _ = T.TONES[emphasis]
            S.rect(slide, _I(BX), _I(y), _I(BW), _I(rh - 0.04), fill=fill,
                   line=line, radius=0.07)
        x = BX
        for j, cell in enumerate(cells):
            w = widths[j] * scale
            colour = T.TEXT if (j == 0 or emphasis) else T.MUTED
            S.textbox(slide, _I(x + 0.14), _I(y + 0.05), _I(w - 0.2),
                      _I(rh - 0.14), str(cell), size=size, color=colour,
                      bold=bool(emphasis) or j == 0, anchor=MSO_ANCHOR.MIDDLE)
            x += w
        if not emphasis and i < len(rows) - 1:
            S.rect(slide, _I(BX), _I(y + rh - 0.04), _I(BW), Pt(0.6),
                   fill=T.LINE, radius=0)
        y += rh

    if spec.get("caption"):
        S.textbox(slide, _I(BX), _I(BY + BH - 0.4), _I(BW), _I(0.4), spec["caption"],
                  size=T.SZ_CAPTION, color=T.WARN, bold=True)


def code(deck, slide, spec):
    deck.heading(slide, spec["title"], spec.get("kicker"))
    lines = spec["code"].rstrip("\n").split("\n")

    # Long listings shrink to fit rather than running off the slide.
    # Rendered line height is about 1.38x the point size at spacing 1.15, so
    # solve the font size against the space available instead of guessing.
    PER_LINE_IN = 0.0192
    PAD = 0.34
    max_h = BH - 0.2
    avail = max_h - PAD
    size = Pt(max(7.5, min(T.SZ_CODE.pt, avail / (len(lines) * PER_LINE_IN))))
    line_h = size.pt * PER_LINE_IN

    h = min(max_h, PAD + line_h * len(lines))
    w = spec.get("width", BW * 0.72)
    S.rect(slide, _I(BX), _I(BY + 0.1), _I(w), _I(h), fill=T.SURFACE, line=T.LINE,
           radius=0.12)
    box = S.textbox(slide, _I(BX + 0.34), _I(BY + 0.24), _I(w - 0.6), _I(h - 0.32),
                    "", size=size, color=T.TEXT)
    tf = box.text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        run = p.add_run()
        run.text = line or " "
        run.font.name = T.MONO
        run.font.size = size
        stripped = line.strip()
        if stripped.startswith("#") or stripped.startswith("//"):
            run.font.color.rgb = T.DIM
        elif spec.get("highlight") and any(k in line for k in spec["highlight"]):
            run.font.color.rgb = T.ACCENT
        else:
            run.font.color.rgb = T.TEXT

    # `notes` is reserved for speaker notes; on-slide annotations are side_notes.
    for i, note in enumerate(spec.get("side_notes", [])):
        ny = BY + 0.22 + i * 1.0
        S.textbox(slide, _I(BX + w + 0.45), _I(ny), _I(BW - w - 0.45), _I(0.9),
                  note, size=T.SZ_CARD_BODY, color=T.MUTED, spacing=1.35)


def takeaways(deck, slide, spec):
    deck.heading(slide, spec.get("title", "Key takeaways"),
                 spec.get("kicker", "Remember this"))
    items = spec["items"]
    h = min(0.95, (BH - 0.6) / len(items) - 0.16)
    for i, item in enumerate(items):
        y = BY + 0.15 + i * (h + 0.16)
        box = S.rect(slide, _I(BX), _I(y), _I(BW), _I(h), fill=T.SURFACE,
                     line=T.LINE, radius=0.1)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Emu(700000)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.2
        S.write(p, item, T.SZ_BODY, T.TEXT)
        chip = S.rect(slide, _I(BX + 0.24), _I(y + (h - 0.34) / 2), _I(0.34),
                      _I(0.34), fill=T.ACCENT, radius=0.08)
        S.label(chip, str(i + 1), size=T.SZ_NODE_SM, color=T.BG, bold=True)


def checklist(deck, slide, spec):
    """Grouped checkbox columns — for deployment and review checklists."""
    deck.heading(slide, spec["title"], spec.get("kicker"))
    groups = spec["groups"]
    gap = 0.34
    cw = (BW - gap * (len(groups) - 1)) / len(groups)
    y0 = BY + 0.12
    text_w = cw - 0.3

    # Items are given the height they actually need. A uniform step collides as
    # soon as one item wraps, and translated items wrap where English did not.
    def _plan(pt):
        line_h = pt * LINE_IN_PER_PT * 1.2
        per_line = max(8, int(text_w / (CHAR_IN_PER_PT * pt)))
        heights, tallest = [], 0.0
        for group in groups:
            hs = [max(line_h + 0.16,
                      math.ceil(len(str(item)) / per_line) * line_h + 0.16)
                  for item in group["items"]]
            heights.append(hs)
            tallest = max(tallest, sum(hs))
        return heights, tallest, line_h

    avail = BH - 0.72
    size = T.SZ_CARD_BODY.pt
    while size > 8.0:
        heights, tallest, _ = _plan(size)
        if tallest <= avail:
            break
        size -= 0.5
    heights, tallest, line_h = _plan(size)

    for i, group in enumerate(groups):
        x = BX + i * (cw + gap)
        S.textbox(slide, _I(x), _I(y0), _I(cw), _I(0.3), group["title"],
                  size=T.SZ_CAPTION, color=T.ACCENT, bold=True, caps=True)
        S.rect(slide, _I(x), _I(y0 + 0.36), _I(cw), Pt(1.2), fill=T.LINE, radius=0)
        iy = y0 + 0.52
        for item, ih in zip(group["items"], heights[i]):
            box = S.rect(slide, _I(x), _I(iy + 0.04), _I(0.15), _I(0.15),
                         fill=None, line=T.DIM, radius=0.02, width=1.1)
            box.text_frame.text = ""
            S.textbox(slide, _I(x + 0.3), _I(iy), _I(text_w), _I(ih),
                      item, size=Pt(size), color=T.TEXT, spacing=1.2)
            iy += ih


def outro(deck, slide, spec):
    S.rect(slide, _I(0), _I(0), _I(0.14), T.SLIDE_H, fill=T.ACCENT, radius=0)
    S.textbox(slide, _I(1.35), _I(2.6), _I(10.2), _I(0.34), spec.get("kicker", ""),
              size=T.SZ_KICKER, color=T.ACCENT, bold=True, caps=True)
    S.textbox(slide, _I(1.35), _I(3.05), _I(10.2), _I(1.6), spec["title"],
              size=T.SZ_DIVIDER_TITLE, color=T.TEXT, bold=True, spacing=1.1)
    if spec.get("note"):
        S.textbox(slide, _I(1.35), _I(4.85), _I(9.6), _I(0.8), spec["note"],
                  size=T.SZ_BODY, color=T.MUTED, spacing=1.3)


RENDERERS = {
    "cover": cover,
    "divider": divider,
    "agenda": agenda,
    "objectives": objectives,
    "bullets": bullets,
    "lead": lead,
    "cards": cards,
    "table": table,
    "code": code,
    "takeaways": takeaways,
    "checklist": checklist,
    "outro": outro,
}
