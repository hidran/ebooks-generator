import pathlib
import re
import sys
import zipfile

import pytest
import yaml
from pptx import Presentation
from pptx.util import Emu

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools/slides"))

import deckgen  # noqa: E402
from deckgen import diagrams, shapes, theme  # noqa: E402

SPEC_DIR = ROOT / "books/neuron-course/slides/en"
SPECS = sorted(SPEC_DIR.glob("module-*.yaml"))

SLIDE_W_IN = 13.333
TOL = 18288  # 0.02" in EMU

_cache = {}


def load(path):
    if path not in _cache:
        spec = yaml.safe_load(path.read_text(encoding="utf-8"))
        _cache[path] = (spec, deckgen.build(spec))
    return _cache[path]


def ids(paths):
    return [p.stem for p in paths]


@pytest.fixture(params=SPECS, ids=ids(SPECS))
def built(request):
    return load(request.param)


# ----------------------------------------------------------------- theme
def test_specs_exist():
    assert SPECS, f"no module specs found in {SPEC_DIR}"


def test_body_area_fits_the_slide():
    assert diagrams.BX + diagrams.BW <= SLIDE_W_IN
    assert diagrams.BY + diagrams.BH <= Emu(theme.FOOTER_Y).inches


def test_body_constants_track_the_theme():
    """diagrams.py caches the body box; it must not drift from theme.py."""
    assert diagrams.BY == pytest.approx(Emu(theme.BODY_Y).inches)
    assert diagrams.BH == pytest.approx(Emu(theme.BODY_H).inches)
    assert diagrams.BW == pytest.approx(Emu(theme.CONTENT_W).inches)


def test_every_tone_resolves_to_three_colours():
    for name, value in theme.TONES.items():
        assert len(value) == 3, name


# ----------------------------------------------------------------- markup
@pytest.mark.parametrize("text,expected", [
    ("plain", [("plain", "plain")]),
    ("call `foo()` now", [("call ", "plain"), ("foo()", "mono"), (" now", "plain")]),
    ("**bold** tail", [("bold", "bold"), (" tail", "plain")]),
    ("", [("", "plain")]),
])
def test_rich_runs_splits_inline_markup(text, expected):
    assert shapes.rich_runs(text) == expected


def test_unknown_layout_is_rejected():
    with pytest.raises(deckgen.SpecError):
        deckgen.build({"deck": {}, "slides": [{"layout": "nope"}]})


# ----------------------------------------------------------------- per deck
def test_slide_count_matches_the_spec(built):
    spec, deck = built
    expected = (
        len(spec.get("slides", []))
        + sum(1 + len(lesson.get("slides", [])) for lesson in spec["lessons"])
        + len(spec.get("closing", []))
    )
    assert deck.slide_count == expected


def test_lessons_open_with_objectives_and_close_with_takeaways(built):
    spec, _ = built
    for lesson in spec["lessons"]:
        layouts = [s["layout"] for s in lesson["slides"]]
        assert layouts[0] == "objectives", lesson["id"]
        assert layouts[-1] == "takeaways", lesson["id"]


def test_all_layouts_used_are_known(built):
    spec, _ = built
    used = {s["layout"] for s in spec.get("slides", [])}
    used |= {s["layout"] for l in spec["lessons"] for s in l["slides"]}
    used |= {s["layout"] for s in spec.get("closing", [])}
    assert used <= set(deckgen.RENDERERS)


def test_shapes_stay_inside_the_slide(built):
    """Nothing may hang off the canvas — it would be cropped on export."""
    _, deck = built
    for i, slide in enumerate(deck.prs.slides):
        for shape in slide.shapes:
            if shape.left is None or shape.width is None:
                continue
            assert shape.left >= -TOL, f"slide {i + 1}: {shape.shape_type} off left"
            assert shape.left + shape.width <= deck.prs.slide_width + TOL, (
                f"slide {i + 1}: {shape.shape_type} overruns the right edge"
            )
            assert shape.top + shape.height <= deck.prs.slide_height + TOL, (
                f"slide {i + 1}: {shape.shape_type} overruns the bottom edge"
            )


def test_every_slide_carries_speaker_notes(built):
    """Decks are recorded from, not just shown — every slide needs a cue."""
    _, deck = built
    missing = [
        i + 1 for i, slide in enumerate(deck.prs.slides)
        if not (slide.has_notes_slide
                and slide.notes_slide.notes_text_frame.text.strip())
    ]
    assert not missing, f"slides without speaker notes: {missing}"


def test_diagram_and_chart_slides_get_a_heading(built):
    """Diagram layouts rely on the engine to draw their title."""
    spec, _ = built
    visual = set(diagrams.RENDERERS) | {"chart", "big_number"}
    for lesson in spec["lessons"]:
        for slide_spec in lesson["slides"]:
            if slide_spec["layout"] in visual:
                assert slide_spec.get("title"), slide_spec["layout"]


def test_code_layout_reserves_notes_for_the_speaker(built):
    """`notes` is speaker notes everywhere; on-slide text is `side_notes`."""
    spec, _ = built
    for lesson in spec["lessons"]:
        for slide_spec in lesson["slides"]:
            if slide_spec["layout"] == "code":
                assert isinstance(slide_spec.get("notes", ""), str)


TRANSLATIONS = [
    (p, SPEC_DIR.parent / lang / p.name)
    for lang in ("it", "es")
    for p in SPECS
    if (SPEC_DIR.parent / lang / p.name).exists()
]


@pytest.mark.parametrize(
    "en_path,tr_path", TRANSLATIONS,
    ids=[f"{t.parent.name}/{t.stem}" for _, t in TRANSLATIONS],
)
def test_translation_matches_english_structure(en_path, tr_path):
    """A translated deck renders the same geometry — only the text changes."""
    en, _ = load(en_path)
    tr = yaml.safe_load(tr_path.read_text(encoding="utf-8"))

    assert [l["id"] for l in tr["lessons"]] == [l["id"] for l in en["lessons"]]

    def layouts(spec):
        return (
            [s["layout"] for s in spec.get("slides", [])]
            + [s["layout"] for l in spec["lessons"] for s in l["slides"]]
            + [s["layout"] for s in spec.get("closing", [])]
        )

    assert layouts(tr) == layouts(en)

    # Code listings and their highlight anchors are never translated.
    for en_lesson, tr_lesson in zip(en["lessons"], tr["lessons"]):
        for a, b in zip(en_lesson["slides"], tr_lesson["slides"]):
            if a["layout"] == "code":
                assert b["code"] == a["code"], f"{tr_path.name} {en_lesson['id']}"
                assert b.get("highlight") == a.get("highlight")


def test_saves_a_readable_pptx(built, tmp_path):
    _, deck = built
    out = tmp_path / "deck.pptx"
    deck.save(out)
    assert out.stat().st_size > 20_000
    assert len(Presentation(str(out)).slides._sldIdLst) == deck.slide_count


def test_notes_master_is_declared_not_just_related(built, tmp_path):
    """Keynote refuses to open a deck whose notesMaster is related from
    presentation.xml but missing from <p:notesMasterIdLst>. python-pptx only
    writes the relationship, so Deck.save has to add the declaration."""
    _, deck = built
    out = tmp_path / "deck.pptx"
    deck.save(out)

    with zipfile.ZipFile(out) as z:
        prs = z.read("ppt/presentation.xml").decode()
        rels = z.read("ppt/_rels/presentation.xml.rels").decode()

    related = re.search(r'Id="([^"]+)"[^>]*/notesMaster"', rels)
    assert related, "these decks all carry speaker notes"

    declared = re.search(
        r"<p:notesMasterIdLst><p:notesMasterId r:id=\"([^\"]+)\"/></p:notesMasterIdLst>",
        prs,
    )
    assert declared, "presentation.xml is missing <p:notesMasterIdLst>"
    assert declared.group(1) == related.group(1)

    # Schema order for CT_Presentation: sldMasterIdLst, notesMasterIdLst, sldIdLst.
    assert (prs.index("sldMasterIdLst")
            < prs.index("notesMasterIdLst")
            < prs.index("sldIdLst"))


def test_slide_size_declares_sixteen_by_nine(built, tmp_path):
    _, deck = built
    out = tmp_path / "deck.pptx"
    deck.save(out)
    with zipfile.ZipFile(out) as z:
        prs = z.read("ppt/presentation.xml").decode()
    assert 'type="screen16x9"' in prs
